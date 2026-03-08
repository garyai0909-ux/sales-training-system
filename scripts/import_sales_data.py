#!/usr/bin/env python3
"""匯入最新銷售數據（Excel）並輸出前端可用的 JSON。"""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl  # type: ignore
import pandas as pd  # type: ignore

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE_DIR = Path(
    "/Users/weiboai/Library/Mobile Documents/com~apple~CloudDocs/WEIBO AI存取區/2026｜銷售數據"
)
DEFAULT_OUTPUT = ROOT / "data" / "dashboard.json"

STORE_HEADERS = {
    "嘉義全店": "WEiZ嘉義旗艦店",
    "高雄全店": "WEiZ高雄旗艦店",
    "台中全店": "WEiZ台中旗艦店",
}
EXCLUDE_KEYS = {"分店/人員", "WEiZ"}


def detect_file(source_dir: Path, month: Optional[str], source_file: Optional[Path]) -> Path:
    if source_file:
        return source_file
    candidates = list(source_dir.glob("20????銷售數據.xlsx"))
    if month:
        for file in candidates:
            if month in file.stem:
                return file
        raise FileNotFoundError(f"在 {source_dir} 找不到 {month} 對應的銷售數據檔")
    if not candidates:
        raise FileNotFoundError(f"{source_dir} 下沒有任何銷售數據檔")
    return max(candidates, key=lambda p: p.stem)


def parse_month_key(path: Path) -> str:
    match = re.search(r"(20\d{4})", path.stem)
    if not match:
        return datetime.now().strftime("%Y%m")
    return match.group(1)


def load_summary(ws) -> Dict[str, Any]:
    headers: Dict[str, int] = {}
    for col in range(1, 40):
        val = ws.cell(row=2, column=col).value
        if val:
            headers[str(val).strip()] = col
    def get_value(row: int, header: str) -> float:
        col = headers.get(header)
        if not col:
            return 0.0
        value = ws.cell(row=row, column=col).value
        if value is None:
            return 0.0
        if isinstance(value, datetime):
            return value.timestamp()
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0
    kpi = {
        "target_revenue": int(get_value(3, "WEiZ")),
        "actual_revenue": int(get_value(4, "WEiZ")),
        "achievement_rate": round(get_value(5, "WEiZ"), 4),
        "total_transactions": int(get_value(6, "WEiZ")),
        "total_pcs": int(get_value(7, "WEiZ")),
    }
    stores: List[Dict[str, Any]] = []
    for header, label in STORE_HEADERS.items():
        target = get_value(3, header)
        actual = get_value(4, header)
        rate = get_value(5, header)
        if not actual:
            continue
        stores.append(
            {
                "name": label,
                "target": int(target),
                "actual": int(actual),
                "achievement_rate": round(rate, 4) if rate else None,
            }
        )
    people: List[Dict[str, Any]] = []
    for name, col in headers.items():
        if name in EXCLUDE_KEYS or not isinstance(name, str):
            continue
        if "店" in name or "全店" in name:
            continue
        amount = ws.cell(row=4, column=col).value
        if not amount:
            continue
        people.append({"name": name, "amount": int(amount)})
    people.sort(key=lambda x: x["amount"], reverse=True)
    return {"kpis": kpi, "stores": stores, "sales_people": people}


def load_top_products(path: Path) -> List[Dict[str, Any]]:
    df = pd.read_excel(path, sheet_name="商品", header=2)
    df = df.rename(columns={"Unnamed: 1": "category", "Unnamed: 2": "product"})
    df = df[df["product"].notna()]
    df = df[df["product"] != "產品名稱"]
    df["WEiZ"] = pd.to_numeric(df["WEiZ"], errors="coerce").fillna(0)
    df = df.sort_values("WEiZ", ascending=False)
    top_rows = df.head(10)
    results = []
    for _, row in top_rows.iterrows():
        results.append(
            {
                "product": str(row["product"]).strip(),
                "category": str(row["category"]).strip(),
                "quantity": int(row["WEiZ"]),
            }
        )
    return results


def build_payload(source_path: Path) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(source_path, data_only=True)
    summary = load_summary(wb["總業績"])
    top_products = load_top_products(source_path)
    return {
        "meta": {
            "source_file": source_path.name,
            "month": parse_month_key(source_path),
            "generated_at": datetime.now().isoformat(),
        },
        "kpis": summary["kpis"],
        "stores": summary["stores"],
        "sales_people": summary["sales_people"],
        "top_products": top_products,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="匯入銷售數據並更新 dashboard JSON")
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR, help="Excel 資料夾")
    parser.add_argument("--source-file", type=Path, help="指定單一 Excel 檔", default=None)
    parser.add_argument("--month", help="指定月份（例如 202602）", default=None)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="輸出的 JSON 路徑")
    args = parser.parse_args()

    source_path = detect_file(args.source_dir, args.month, args.source_file)
    payload = build_payload(source_path)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2))
    print(f"已將 {source_path.name} 匯入並輸出到 {args.output}")


if __name__ == "__main__":
    main()
